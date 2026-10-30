"""Human validation (brief: 100-150 tweets, 2 coders, Cohen's kappa >= 0.7).

1. `sample` draws a stratified sample and writes one blind coding sheet per coder
   (the engine's labels are not shown, so coders are not anchored by them).
2. `score` compares the two coders, builds the gold labels (agreed values, or an
   adjudicated sheet), and measures the engine against them, per field and per keyword.
"""

import json
import math
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation
from sklearn.metrics import classification_report, cohen_kappa_score

from xetl.engine import band
from xetl.stats import md_table

FIELDS = ["sentiment_score", "legitimacy_cognitive", "legitimacy_pragmatic", "legitimacy_moral"]
LEGITIMACY_VALUES = ["Granted", "Denied", "NA"]
KAPPA_TARGET = 0.7

GUIDE = [
    ("sentiment_score", "1-2 negative (1 strong), 3 neutral, 4-5 positive (5 strong)."),
    ("legitimacy_cognitive", "Recognized as a real business (Granted) or as illegitimate (Denied)."),
    ("legitimacy_pragmatic", "Worth the exchange: good product/price (Granted) or not (Denied)."),
    ("legitimacy_moral", "Ethically deserving support (Granted) or condemned on principle (Denied)."),
    ("", "Use NA when a legitimacy type does not apply. A tweet can have several or none."),
    ("sarcasm", "Y if the literal wording means the opposite (e.g. 'wah bagus lah ... tak payah lesen')."),
    ("", "Code independently. Do not discuss tweets with the other coder until both are done."),
]  # fmt: skip


def draw_sample(coded: pd.DataFrame, n: int, seed: int) -> pd.DataFrame:
    """Equal share per sentiment band where possible, topped up at random, then shuffled."""
    n = min(n, len(coded))
    per_band = math.ceil(n / 3)
    picked = coded.groupby("sentiment_band", group_keys=False).apply(
        lambda g: g.sample(min(len(g), per_band), random_state=seed)
    )
    if len(picked) < n:
        rest = coded.drop(picked.index)
        picked = pd.concat([picked, rest.sample(n - len(picked), random_state=seed)])
    return picked.head(n).sample(frac=1, random_state=seed)


def write_coding_sheet(sample: pd.DataFrame, path: Path, prefill: pd.DataFrame | None = None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Coding"
    headers = ["tweet_id", "text", *FIELDS, "sarcasm", "notes"]
    extra = [] if prefill is None else [c for c in prefill.columns if c.startswith("coder")]
    ws.append(headers + extra)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for i, row in enumerate(sample.itertuples(index=False)):
        values = [row.tweet_id, row.text_clean, None, None, None, None, None, None]
        if prefill is not None:
            p = prefill.iloc[i]
            values[2:6] = [None if pd.isna(p[f]) else p[f] for f in FIELDS]
            values += [p[c] for c in extra]
        ws.append(values)

    last = ws.max_row
    score_dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5")
    legit_dv = DataValidation(type="list", formula1='"Granted,Denied,NA"')
    yn_dv = DataValidation(type="list", formula1='"Y,N"')
    for dv, cols in ((score_dv, "C"), (legit_dv, "DEF"), (yn_dv, "G")):
        ws.add_data_validation(dv)
        for col in cols:
            dv.add(f"{col}2:{col}{last}")
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 80
    for col in "CDEFGH":
        ws.column_dimensions[col].width = 20
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"

    guide = wb.create_sheet("Guide")
    for row in GUIDE:
        guide.append(row)
    guide.column_dimensions["A"].width = 22
    guide.column_dimensions["B"].width = 100
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def read_coding_sheet(path: Path) -> pd.DataFrame:
    # keep_default_na=False: pandas would otherwise read the label "NA" as an empty cell.
    df = pd.read_excel(
        path, sheet_name="Coding", dtype={"tweet_id": str}, keep_default_na=False, na_values=[""]
    )
    df = df[["tweet_id", *FIELDS]].copy()
    missing = int(df[FIELDS].isna().sum().sum())
    if missing:
        raise ValueError(f"{path.name}: {missing} empty cells in {FIELDS}; finish coding first")
    df["sentiment_score"] = df["sentiment_score"].astype(int)
    for f in FIELDS[1:]:
        df[f] = df[f].astype(str).str.strip().str.capitalize().replace({"Na": "NA"})
        bad = set(df[f]) - set(LEGITIMACY_VALUES)
        if bad:
            raise ValueError(f"{path.name}: unexpected values in {f}: {sorted(bad)}")
    return df.set_index("tweet_id")


def kappa(a: pd.Series, b: pd.Series, weights: str | None = None) -> float | None:
    """Cohen's kappa, or None when it is undefined (only one label used overall)."""
    if len(set(a) | set(b)) < 2:
        return None
    return round(float(cohen_kappa_score(a, b, weights=weights)), 3)


def _fmt(k: float | None) -> str:
    if k is None:
        return "n/a"
    return f"{k:.3f} ({'meets' if k >= KAPPA_TARGET else 'below'} {KAPPA_TARGET})"


def keyword_precision(coded: pd.DataFrame, gold: pd.DataFrame) -> pd.DataFrame:
    """For each matched keyword: how often its label agreed with the gold label."""
    rows = []
    for tweet_id, matches_json in coded["matches_json"].items():
        if tweet_id not in gold.index:
            continue
        for m in json.loads(matches_json):
            if m["dimension"] == "sentiment":
                truth = gold.at[tweet_id, "sentiment_score"]
                correct = pd.notna(truth) and band(int(m["label"])) == band(int(truth))
            else:
                truth = gold.at[tweet_id, f"legitimacy_{m['dimension']}"]
                correct = pd.notna(truth) and m["label"] == truth
            if pd.notna(truth):
                rows.append((m["dimension"], m["keyword"], correct))
    if not rows:
        return pd.DataFrame(columns=["dimension", "keyword", "matches", "precision"])
    df = pd.DataFrame(rows, columns=["dimension", "keyword", "correct"])
    out = df.groupby(["dimension", "keyword"])["correct"].agg(matches="size", precision="mean")
    out["precision"] = out["precision"].round(2)
    return out.reset_index().sort_values(["precision", "matches"], ascending=[True, False])


def score_validation(
    coded: pd.DataFrame, coder1: Path, coder2: Path, out_dir: Path, adjudicated: Path | None
) -> Path:
    c1, c2 = read_coding_sheet(coder1), read_coding_sheet(coder2)
    ids = c1.index.intersection(c2.index)
    c1, c2 = c1.loc[ids], c2.loc[ids]
    coded = coded.set_index("tweet_id").loc[ids]

    inter = pd.DataFrame(
        [
            {
                "field": f,
                "kappa": _fmt(kappa(c1[f], c2[f], "quadratic" if f == "sentiment_score" else None)),
                "% agreement": round(100 * (c1[f] == c2[f]).mean(), 1),
            }
            for f in FIELDS
        ]
    )
    bands1, bands2 = c1["sentiment_score"].map(band), c2["sentiment_score"].map(band)
    band_row = {"field": "sentiment_band", "kappa": _fmt(kappa(bands1, bands2)),
                "% agreement": round(100 * (bands1 == bands2).mean(), 1)}  # fmt: skip
    inter = pd.concat([inter, pd.DataFrame([band_row])], ignore_index=True)

    if adjudicated:
        gold = read_coding_sheet(adjudicated).reindex(ids)
        gold_source = f"adjudicated sheet `{adjudicated.name}`"
    else:
        gold = c1.where(c1 == c2)  # agreed values only; disagreements stay empty
        gold_source = "values both coders agreed on (disagreements excluded)"

    engine_rows, class_tables = [], []
    for f in FIELDS:
        mask = gold[f].notna()
        truth, pred = gold.loc[mask, f], coded.loc[mask, f]
        if f == "sentiment_score":
            truth, pred = truth.astype(int), pred.astype(int)
        engine_rows.append(
            {
                "field": f,
                "n": int(mask.sum()),
                "kappa": _fmt(kappa(truth, pred, "quadratic" if f == "sentiment_score" else None)),
                "accuracy": round(float((truth == pred).mean()), 3) if mask.any() else "n/a",
            }
        )
        if mask.any():
            report = classification_report(truth, pred, output_dict=True, zero_division=0)
            table = pd.DataFrame(
                [
                    {"label": k, "precision": round(v["precision"], 2),
                     "recall": round(v["recall"], 2), "support": int(v["support"])}  # fmt: skip
                    for k, v in report.items()
                    if isinstance(v, dict) and k not in ("macro avg", "weighted avg")
                ]
            )
            class_tables += [f"\n### {f}\n", md_table(table)]

    disagree = (c1 != c2).any(axis=1)
    prefill = c1.where(c1 == c2)
    for f in FIELDS:
        prefill[f"coder1_{f}"], prefill[f"coder2_{f}"] = c1[f], c2[f]
    write_coding_sheet(
        coded.reset_index(),
        out_dir / "adjudication.xlsx",
        prefill=prefill.reset_index(drop=True),
    )

    precision = keyword_precision(coded, gold)
    report_lines = [
        "# Validation report",
        f"- Tweets coded by both coders: **{len(ids)}**",
        f"- Tweets where coders disagree on at least one field: **{int(disagree.sum())}** "
        "(resolve them in `adjudication.xlsx`, then re-run with `--adjudicated`)",
        f"- Gold labels: {gold_source}",
        "\n## Inter-coder agreement\n",
        "Sentiment score uses quadratic-weighted kappa (1 vs 2 is a smaller disagreement "
        "than 1 vs 5).\n",
        md_table(inter),
        "\n## Engine vs gold\n",
        md_table(pd.DataFrame(engine_rows)),
        "\n## Per-class precision and recall",
        *class_tables,
        "\n## Keyword precision (lowest first: candidates for codebook v2)\n",
        md_table(precision),
    ]
    out = out_dir / "validation_report.md"
    out.write_text("\n".join(report_lines) + "\n", encoding="utf-8")
    return out
