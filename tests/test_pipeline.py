"""End-to-end on fake extracted data: transform -> analyze -> stats -> export -> validation."""

import json

import pandas as pd
import pytest
from conftest import CONFIG, TEMPLATE
from openpyxl import load_workbook

from xetl.analyze import run_analyze
from xetl.config import Settings
from xetl.export import run_export
from xetl.stats import run_stats
from xetl.transform import run_transform
from xetl.validation import FIELDS, draw_sample, score_validation, write_coding_sheet

POSTS = [
    ("1", "Kedai rohingya ni murah dan sedap, saya sokong #RohingyaMalaysia", None, None),
    ("2", "Macam mana boleh jual kalau tak berdaftar? #rohingyamalaysia", None, None),
    ("3", "RT @someone: Kedai ni sedap #RohingyaMalaysia", "99", None),
    ("4", "Kasihan diorang, tiada pilihan, patut dibantu #Rohingya_Malaysia", None, None),
    ("5", "=HYPERLINK(\"http://evil\") penipu #RohingyaMalaysia", None, None),
    ("6", "Cuaca panas hari ini #Rohingya #Malaysia", None, None),
    ("7", "Penipu! Busuk punya kedai, jangan beli #RohingyaMalaysia", None, None),
    ("1", "duplicate id is dropped #RohingyaMalaysia", None, None),
]


@pytest.fixture
def settings(tmp_path, monkeypatch):
    monkeypatch.setenv("XETL_AUTHOR_HASH_SALT", "test-salt")
    s = Settings(data_dir=tmp_path / "data", config_dir=CONFIG, template_path=TEMPLATE)
    raw = s.stage_dir("raw", "t")
    raw.mkdir(parents=True)
    with (raw / "2026-01-05.jsonl").open("w") as f:
        for tweet_id, text, retweeted, quoted in POSTS:
            f.write(json.dumps({
                "tweet_id": tweet_id, "timestamp": "2026-01-05T03:00:00+00:00", "text": text,
                "hashtags": [], "lang": "in", "like_count": 2, "retweet_count": 1,
                "reply_count": 0, "quote_count": 0, "view_count": 50, "author_id_hash": "abc",
                "conversation_id": tweet_id, "in_reply_to_id": None, "retweeted_id": retweeted,
                "quoted_id": quoted,
            }) + "\n")  # fmt: skip
    return s


def test_pipeline(settings, hashtags, engine, rules):
    tweets = run_transform(settings, hashtags, "t")
    assert len(tweets) == 7  # duplicate ID removed
    assert tweets.set_index("tweet_id").at["6", "hashtag_match"] == "split_pair"
    assert tweets.set_index("tweet_id").at["3", "post_type"] == "retweet"

    coded = run_analyze(settings, engine, "t")
    assert set(coded["tweet_id"]) == {"1", "2", "4", "5", "7"}  # no retweet, no split pair

    report = run_stats(settings, hashtags, "t").read_text()
    assert "Posts collected: **7**" in report
    assert "#rohingyamalaysia" in report

    if TEMPLATE.exists():
        out = run_export(settings, rules, hashtags, "t")
        wb = load_workbook(out)
        assert len(wb["Sentiment_Summary"]._charts) == 2
        assert len(wb["Legitimacy_Summary"]._charts) == 1
        assert "$C$2:$C$6" in wb["Sentiment_Summary"]["B5"].value
        assert wb["Coded_Data"].max_row == 6
        assert wb["Coded_Data"]["D2"].value.startswith("=IF(C2")
        evil = [c for c in wb["Raw_Data"]["C"] if str(c.value).startswith("=HYPERLINK")]
        assert evil and evil[0].data_type == "s"  # stored as text, not a formula
        assert "Engine_Trace" in wb.sheetnames

    sample = draw_sample(coded, n=5, seed=1)
    out_dir = settings.stage_dir("outputs", "t") / "validation"
    for coder in ("coder1", "coder2"):
        path = out_dir / f"coding_sheet_{coder}.xlsx"
        write_coding_sheet(sample, path)
        filled = pd.read_excel(path, sheet_name="Coding", dtype={"tweet_id": str})
        truth = coded.set_index("tweet_id").loc[filled["tweet_id"], FIELDS].reset_index(drop=True)
        filled[FIELDS] = truth
        if coder == "coder2":
            # one disagreement
            filled.loc[0, "sentiment_score"] = 1 if truth.loc[0, "sentiment_score"] != 1 else 5
        with pd.ExcelWriter(path, mode="a", if_sheet_exists="replace") as writer:
            filled.to_excel(writer, sheet_name="Coding", index=False)

    report = score_validation(
        coded, out_dir / "coding_sheet_coder1.xlsx", out_dir / "coding_sheet_coder2.xlsx",
        out_dir, adjudicated=None,
    ).read_text()  # fmt: skip
    assert "Tweets coded by both coders: **5**" in report
    assert "disagree on at least one field: **1**" in report
    assert (out_dir / "adjudication.xlsx").exists()
